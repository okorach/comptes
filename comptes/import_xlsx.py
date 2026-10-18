"""One-time import of Comptes.xlsm into the SQLite database.

Usage: python -m comptes.import_xlsx "D:\\Backup\\GDrive\\Banque\\Comptes.xlsm" [comptes.db]

Values are read from the workbook's cached results (data_only), so history is
migrated exactly as Excel computed it. Formulas are inspected only to detect how
each account is maintained (mode: flow / snapshot / shares / static).
"""
import sys, os, datetime as dt
import openpyxl
from .db import connect, DEFAULT_DB

UNDATED = []
HDR_ROW = 10
FIRST = 11


def iso(d):
    return d.date().isoformat() if isinstance(d, dt.datetime) else d.isoformat()


def num(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def read_params(wbv, con):
    ws = wbv["Paramètres"]
    for r in ws.iter_rows(min_row=2, max_row=200, max_col=18, values_only=True):
        sub, cat = r[0], r[1]
        if sub and cat:
            con.execute("INSERT OR REPLACE INTO categories VALUES(?,?)", (sub, cat))
        if r[12] and r[13]:  # M,N substitutions
            con.execute("INSERT INTO substitutions(pattern,replacement) VALUES(?,?)", (r[12], r[13]))
        if r[15] and r[16] is not None:  # P,Q,R account types
            con.execute("INSERT OR REPLACE INTO account_types VALUES(?,?,?)", (r[15], int(r[16]), num(r[17]) or 0))
    con.execute("INSERT OR REPLACE INTO params VALUES('default_currency','EUR')")


def read_fx(wbv, con):
    ws = wbv["Exchange Rates"]
    for r in ws.iter_rows(min_row=3, max_col=8, values_only=True):
        for cur, di, ri in (("CHF", 0, 1), ("USD", 3, 4), ("GBP", 6, 7)):
            if isinstance(r[di], dt.datetime) and num(r[ri]):
                con.execute("INSERT OR REPLACE INTO fx_rates VALUES(?,?,?)", (cur, iso(r[di]), float(r[ri])))


def read_accounts(wbv, con):
    ws = wbv["Comptes"]
    names = []
    for r in ws.iter_rows(min_row=3, max_col=11, values_only=True):
        if not r[2]:
            continue
        status = r[5]
        con.execute(
            "INSERT OR REPLACE INTO accounts(id,number,bank,availability,status,currency,type,in_budget,tax,website)"
            " VALUES(?,?,?,?,?,?,?,?,?,?)",
            (r[2], None if r[1] is None else str(r[1]), r[3], int(r[4] or 0), status, r[6], r[7],
             1 if r[8] else 0, num(r[9]) or 0, r[10]))
        names.append(r[2])
    return names


def header_map(row):
    """Map header name -> first column index; also locate the deposits table (2nd 'Date')."""
    cols, date_idx = {}, []
    for i, h in enumerate(row):
        if h is None:
            continue
        if h == "Date":
            date_idx.append(i)
        cols.setdefault(h, i)
    return cols, date_idx


def detect_mode(cols, frow, cur):
    if "Nbr of shares" in cols:
        return "shares"
    amt_col = cols.get("Montant " + cur, cols.get("Montant")) if cur != "EUR" else cols.get("Montant")
    f = frow[amt_col] if frow and amt_col is not None and amt_col < len(frow) else None
    if isinstance(f, str) and f.startswith("="):
        if "Solde" in f and "Francs" not in f:
            return "snapshot"
        return "static"
    return "flow"


def import_account(name, wbv, wbf, con):
    cur = con.execute("SELECT currency FROM accounts WHERE id=?", (name,)).fetchone()["currency"]
    wsv, wsf = wbv[name], wbf[name]
    vrows = list(wsv.iter_rows(min_row=HDR_ROW, max_col=16, values_only=True))
    if not vrows:
        return 0, 0
    cols, date_idx = header_map(vrows[0])
    dep_col = date_idx[1] if len(date_idx) > 1 else None
    limit = dep_col if dep_col is not None else 99
    frow = next(wsf.iter_rows(min_row=FIRST, max_row=FIRST, max_col=16, values_only=True), None)
    mode = detect_mode(cols, frow, cur)
    con.execute("UPDATE accounts SET mode=? WHERE id=?", (mode, name))
    c = lambda h: (cols[h] if h in cols and cols[h] < limit else None)
    i_amt, i_bal = c("Montant"), c("Solde")
    i_nat_amt = c("Montant " + cur) if cur != "EUR" else None
    i_nat_bal = c("Solde " + cur) if cur != "EUR" else None
    i_desc, i_sub, i_bud = c("Description"), c("Sous-Catégorie"), c("In Budget")
    i_qty, i_px = c("Nbr of shares"), c("Price per share")
    get = lambda r, i: None if i is None or i >= len(r) else r[i]
    n = nd = 0
    prev = None
    for r in vrows[1:]:
        d = r[0]
        if not isinstance(d, dt.datetime) and prev is not None and (isinstance(get(r, i_amt), (int, float)) or isinstance(get(r, i_bal), (int, float))) and get(r, i_desc):
            # Row typed without a date in the workbook: keep it (the workbook counts it in the balance),
            # dated like the previous row so that nothing is lost.
            d = prev
            UNDATED.append((name, iso(d), get(r, i_desc)))
        if isinstance(d, dt.datetime):
            prev = d
            a_eur, b_eur = num(get(r, i_amt)), num(get(r, i_bal))
            if cur == "EUR":
                a_nat, b_nat = a_eur, b_eur
            else:
                a_nat, b_nat = num(get(r, i_nat_amt)), num(get(r, i_nat_bal))
            con.execute(
                "INSERT INTO transactions(account,date,amount_eur,balance_eur,native_amount,native_balance,"
                "description,subcategory,in_budget,qty,price) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                (name, iso(d), a_eur, b_eur, a_nat, b_nat, get(r, i_desc), get(r, i_sub),
                 num(get(r, i_bud)), num(get(r, i_qty)), num(get(r, i_px))))
            n += 1
        if dep_col is not None and isinstance(get(r, dep_col), dt.datetime):
            a = num(get(r, dep_col + 1))
            nat = num(get(r, dep_col + 2)) if cur != "EUR" and "Montant " + cur in cols else a
            if a is not None:
                con.execute("INSERT INTO deposits(account,date,amount_eur,native_amount) VALUES(?,?,?,?)",
                            (name, iso(get(r, dep_col)), a, nat))
                nd += 1
    return n, nd


def read_interests(wbv, con):
    ws = wbv["Interests"]
    for r in ws.iter_rows(min_row=2, max_row=200, max_col=7, values_only=True):
        if r[0] and isinstance(r[1], dt.datetime):
            v = [num(x) for x in r[2:7]]
            con.execute("INSERT OR REPLACE INTO interests VALUES(?,?,?,?,?,?,?)",
                        (r[0], r[1].isoformat(), *v))


def main(xlsm, dbpath=None):
    dbpath = dbpath or DEFAULT_DB
    if os.path.exists(dbpath):
        os.remove(dbpath)
    con = connect(dbpath)
    wbv = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    wbf = openpyxl.load_workbook(xlsm, read_only=True, data_only=False)
    read_params(wbv, con)
    read_fx(wbv, con)
    names = read_accounts(wbv, con)
    for name in names:
        if name not in wbv.sheetnames:
            print("no sheet for", name)
            continue
        n, nd = import_account(name, wbv, wbf, con)
        print(f"{name}: {n} transactions, {nd} deposits", flush=True)
    read_interests(wbv, con)
    con.commit()
    if UNDATED:
        print(f"WARNING: {len(UNDATED)} rows had no date in the workbook; they were dated like the previous row:")
        for u in UNDATED:
            print("   ", u)
    print("done ->", dbpath)


if __name__ == "__main__":
    main(*sys.argv[1:3])
